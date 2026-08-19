# -*- coding: utf-8 -*-
"""영수증 정산 세부내역 Excel 내보내기 (openpyxl).

확정된 card 목록을 한 시트에 행으로 쌓는다. 미확인 금액(None)은 빈 셀로 두고
0원과 구분한다. 합계 행은 숫자로 쓰인 셀만 더하고, 더할 값이 없으면 합계도 빈 셀이다.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET_NAME = "영수증정산"
HEADERS = [
    "번호", "파일명", "상호", "사업자번호", "거래일시",
    "공급가액", "부가세", "합계", "카드번호", "승인번호", "품목", "검증",
]
COL_SUPPLY = 6
COL_VAT = 7
COL_TOTAL = 8
COL_ITEMS = 11
COL_FLAGS = 12
AMOUNT_FORMAT = "#,##0"
FONT_NAME = "맑은 고딕"
# 번호, 파일명, 상호, 사업자번호, 거래일시, 공급가액, 부가세, 합계, 카드번호, 승인번호, 품목, 검증
COL_WIDTHS = (6, 22, 20, 16, 18, 12, 12, 12, 22, 14, 36, 36)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")


def _is_number(v):
    """합산 대상. bool 은 int 하위라 명시적으로 제외. None/빈값은 숫자 아님."""
    if v is None or isinstance(v, bool):
        return False
    return isinstance(v, (int, float))


def _amount_value(v):
    """금액 셀에 쓸 값. None/빈문자열은 None(빈 셀). 실제 0은 0."""
    if v is None or v == "":
        return None
    if not _is_number(v):
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return int(v) if isinstance(v, int) else v


def _item_line(it):
    """한 품목: '이름 x수량 금액'. 수량 None 이면 x수량 표기를 생략."""
    if not isinstance(it, dict):
        return ""
    name = str(it.get("name") or "").strip()
    qty = it.get("qty")
    price = it.get("price")
    parts = []
    if name:
        parts.append(name)
    if qty is not None:
        parts.append(f"x{qty}")
    if price is not None:
        if _is_number(price):
            n = int(price) if isinstance(price, float) and price.is_integer() else price
            parts.append(str(int(n) if isinstance(n, int) else n))
        else:
            parts.append(str(price))
    return " ".join(parts)


def _items_text(items):
    lines = [_item_line(it) for it in (items or [])]
    return "\n".join(line for line in lines if line)


def _flags_text(card):
    texts = [str(t) for t in (card.get("flag_texts") or []) if t]
    return ", ".join(texts) if texts else "정상"


def _write_amount(cell, value, font, align):
    cell.font = font
    cell.alignment = align
    cell.number_format = AMOUNT_FORMAT
    n = _amount_value(value)
    if n is not None:
        cell.value = n
        return n
    return None


def export_receipts(items: list, out_path: str) -> dict:
    """확정된 card 목록을 정산 세부내역 시트로 저장. 반환 {"ok":bool,"path":str,"count":int,"error":str}"""
    path = str(out_path or "")
    try:
        if not path:
            return {"ok": False, "path": "", "count": 0, "error": "저장 경로가 없습니다."}

        font = Font(name=FONT_NAME, size=10)
        header_font = Font(name=FONT_NAME, size=10, bold=True)
        total_font = Font(name=FONT_NAME, size=10, bold=True)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        text_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        amount_align = Alignment(horizontal="right", vertical="top")
        center_align = Alignment(horizontal="center", vertical="top")

        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

        for col, header in enumerate(HEADERS, 1):
            cell = ws.cell(1, col, header)
            cell.font = header_font
            cell.fill = HEADER_FILL
            cell.alignment = header_align
        for col, width in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 18

        # 합산 대상 숫자 셀이 하나도 없으면 합계도 빈 셀로 둔다 — 전부 미확인인 열에
        # 0을 찍으면 "공급가액 합계 0원"이라는 없는 사실이 정산서에 생긴다.
        supply_sum = vat_sum = total_sum = None
        count = 0
        for card in (items or []):
            if not isinstance(card, dict):
                continue
            data = card.get("data") if isinstance(card.get("data"), dict) else {}
            count += 1
            row = count + 1

            ws.cell(row, 1, count).font = font
            ws.cell(row, 1).alignment = center_align
            ws.cell(row, 2, str(card.get("name") or "")).font = font
            ws.cell(row, 2).alignment = text_align
            ws.cell(row, 3, str(data.get("store_name") or "")).font = font
            ws.cell(row, 3).alignment = text_align
            ws.cell(row, 4, str(data.get("biz_no") or "")).font = font
            ws.cell(row, 4).alignment = text_align
            ws.cell(row, 5, str(data.get("tx_datetime") or "")).font = font
            ws.cell(row, 5).alignment = text_align

            n = _write_amount(ws.cell(row, COL_SUPPLY), data.get("supply_amount"),
                              font, amount_align)
            if n is not None:
                supply_sum = (supply_sum or 0) + n
            n = _write_amount(ws.cell(row, COL_VAT), data.get("vat"),
                              font, amount_align)
            if n is not None:
                vat_sum = (vat_sum or 0) + n
            n = _write_amount(ws.cell(row, COL_TOTAL), data.get("total_amount"),
                              font, amount_align)
            if n is not None:
                total_sum = (total_sum or 0) + n

            ws.cell(row, 9, str(data.get("card_no_masked") or "")).font = font
            ws.cell(row, 9).alignment = text_align
            ws.cell(row, 10, str(data.get("approval_no") or "")).font = font
            ws.cell(row, 10).alignment = text_align
            item_cell = ws.cell(row, COL_ITEMS, _items_text(data.get("items")))
            item_cell.font = font
            item_cell.alignment = text_align
            flag_cell = ws.cell(row, COL_FLAGS, _flags_text(card))
            flag_cell.font = font
            flag_cell.alignment = text_align

        total_row = count + 2
        label = ws.cell(total_row, 1, "합계")
        label.font = total_font
        label.alignment = center_align
        for col in range(2, 13):
            ws.cell(total_row, col).font = font
        _write_amount(ws.cell(total_row, COL_SUPPLY), supply_sum, total_font, amount_align)
        _write_amount(ws.cell(total_row, COL_VAT), vat_sum, total_font, amount_align)
        _write_amount(ws.cell(total_row, COL_TOTAL), total_sum, total_font, amount_align)

        wb.save(path)
        return {"ok": True, "path": path, "count": count, "error": ""}
    except Exception as e:
        return {"ok": False, "path": path, "count": 0, "error": str(e)}
